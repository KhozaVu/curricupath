"""Derive timetable-clash rules from the two School registration workbooks."""

from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RAW = PROJECT_ROOT / "data" / "raw"
RULES = PROJECT_ROOT / "data" / "rules"
PLANNER_PATH = RAW / "Student_Curricular_Planner_2026.xlsx"
CROSS_CURRICULAR_PATH = RAW / "Cross_Curricular_Registration_-_Requirement_and_Clashes.xlsx"
GRID_RULES_PATH = RULES / "grid_rules.json"
OUTPUT_PATH = RULES / "grid_clash_rules.csv"
DISAGREEMENTS_PATH = RULES / "clash_source_disagreements.csv"
TEMPLATE_PATH = RULES / "clash_cases_template.csv"

PLANNER_SOURCE = "EIE_PLANNER_2026"
TEXT_SOURCE = "EIE_REG_GRID_2026"
BLUE_FILL = "FF00B0F0"
CODE_PATTERN = re.compile(r"([A-Z]{4})\s*(\d{4})([A-Z]?)", re.IGNORECASE)
CLASH_PATTERN = re.compile(
    r"([A-Z]{4}\s*\d{4}[A-Z]?)(?:\s*\((LEC|LAB|TUT)\))?",
    re.IGNORECASE,
)
LEGACY_CODE_MAP = {
    "ELEN2004": "ELEN2020",
    "ELEN2001": "ELEN2016",
    "ELEN2006": "ELEN2021",
    "MATH1048": "MATH1042",
    "MATH1049": "MATH1043",
}
IGNORED_CODES = {"MATH1014"}
PLANNER_SHEETS = (
    ("EE YOS 2", 6, "EFA03"),
    ("Elec YoS 3", 4, "EFA03"),
    ("Info YoS 3", 4, "EFA04"),
)
TEXT_SHEETS = (
    ("EE YOS 2", "EFA03"),
    ("EE YOS 3", "EFA03"),
    ("IE YOS 3", "EFA04"),
)
CLASH_TYPES = ("LEC", "LAB", "TUT", "UNSPECIFIED")


@dataclass(frozen=True)
class SourceClash:
    scope: str
    repeating_course: str
    target_course_code: str
    clash_type: str
    source_id: str
    sheet: str
    raw_value: str


def canonical_code(value: object) -> str | None:
    """Return a current code with an A suffix, applying documented legacy mappings."""

    match = CODE_PATTERN.search(str(value or "").upper())
    if match is None:
        return None
    base = f"{match.group(1).upper()}{match.group(2)}"
    if base in IGNORED_CODES:
        return None
    base = LEGACY_CODE_MAP.get(base, base)
    return f"{base}{match.group(3).upper() or 'A'}"


def _base(code: str) -> str:
    return code[:-1] if re.fullmatch(r"[A-Z]{4}\d{4}[A-Z]", code) else code


def _fill_rgb(cell) -> str:
    colour = cell.fill.fgColor
    if cell.fill.fill_type is None or colour.type != "rgb" or not colour.rgb:
        return ""
    rgb = str(colour.rgb).upper()
    return f"FF{rgb}" if len(rgb) == 6 else rgb


def _clash_type(value: object) -> str:
    text = str(value or "").upper()
    for clash_type in ("LEC", "LAB", "TUT"):
        if clash_type in text:
            return clash_type
    return "UNSPECIFIED"


def extract_planner_clashes(path: Path = PLANNER_PATH) -> list[SourceClash]:
    workbook = load_workbook(path, data_only=False)
    clashes: list[SourceClash] = []
    for sheet_name, header_row, scope in PLANNER_SHEETS:
        sheet = workbook[sheet_name]
        targets = {
            column: canonical_code(sheet.cell(header_row, column).value)
            for column in range(2, sheet.max_column + 1)
        }
        for row in range(header_row + 1, sheet.max_row + 1):
            repeating = canonical_code(sheet.cell(row, 1).value)
            if repeating is None:
                continue
            for column, target in targets.items():
                cell = sheet.cell(row, column)
                if target is None or _fill_rgb(cell) != BLUE_FILL:
                    continue
                clashes.append(
                    SourceClash(
                        scope=scope,
                        repeating_course=repeating,
                        target_course_code=target,
                        clash_type=_clash_type(cell.value),
                        source_id=PLANNER_SOURCE,
                        sheet=sheet_name,
                        raw_value=str(cell.value or ""),
                    )
                )
    return clashes


def assert_planner_matches_grid_rules(
    clashes: Iterable[SourceClash], path: Path = GRID_RULES_PATH
) -> None:
    """Fail if workbook colour decoding differs from the supplied decoded JSON."""

    supplied = json.loads(path.read_text(encoding="utf-8"))["planner_grid_pairs"]
    expected = {
        (
            row["sheet"],
            canonical_code(row["repeating"]),
            canonical_code(row["target"]),
            _clash_type(row.get("detail")),
        )
        for row in supplied
        if row.get("kind") == "timetable_clash"
    }
    actual = {
        (row.sheet, row.repeating_course, row.target_course_code, row.clash_type)
        for row in clashes
    }
    if actual != expected:
        added = sorted(actual - expected)
        missing = sorted(expected - actual)
        raise ValueError(
            "Planner colour extraction differs from grid_rules.json: "
            f"added={added[:5]}, missing={missing[:5]}."
        )


def extract_text_clashes(path: Path = CROSS_CURRICULAR_PATH) -> list[SourceClash]:
    workbook = load_workbook(path, data_only=False)
    clashes: list[SourceClash] = []
    for sheet_name, scope in TEXT_SHEETS:
        sheet = workbook[sheet_name]
        header_row = next(
            row
            for row in range(1, min(sheet.max_row, 10) + 1)
            if str(sheet.cell(row, 1).value or "").strip().lower() == "courses"
        )
        headers = {
            str(sheet.cell(header_row, column).value or "").strip(): column
            for column in range(1, sheet.max_column + 1)
        }
        clash_column = headers["Time-table Clashes"]
        for row in range(header_row + 1, sheet.max_row + 1):
            repeating = canonical_code(sheet.cell(row, 1).value)
            raw = str(sheet.cell(row, clash_column).value or "").strip()
            if repeating is None or not raw or raw.upper() == "NONE":
                continue
            for match in CLASH_PATTERN.finditer(raw):
                target = canonical_code(match.group(1))
                if target is None or target.startswith(("BME", "DIGA", "WSOA")):
                    continue
                clashes.append(
                    SourceClash(
                        scope=scope,
                        repeating_course=repeating,
                        target_course_code=target,
                        clash_type=(match.group(2) or "UNSPECIFIED").upper(),
                        source_id=TEXT_SOURCE,
                        sheet=sheet_name,
                        raw_value=raw,
                    )
                )
    return clashes


def _selected_type(rows: list[SourceClash]) -> str:
    types = {row.clash_type for row in rows}
    for clash_type in CLASH_TYPES:
        if clash_type in types:
            return clash_type
    return "UNSPECIFIED"


def combine_clashes(
    planner: Iterable[SourceClash], text: Iterable[SourceClash]
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    grouped: dict[tuple[str, str, str], list[SourceClash]] = defaultdict(list)
    for row in [*planner, *text]:
        grouped[(row.scope, row.repeating_course, row.target_course_code)].append(row)

    rules: list[dict[str, str]] = []
    disagreements: list[dict[str, str]] = []
    for (scope, repeating, target), rows in sorted(grouped.items()):
        planner_rows = [row for row in rows if row.source_id == PLANNER_SOURCE]
        text_rows = [row for row in rows if row.source_id == TEXT_SOURCE]
        agreement = (
            "both" if planner_rows and text_rows else "planner_only" if planner_rows else "text_only"
        )
        planner_types = sorted({row.clash_type for row in planner_rows})
        text_types = sorted({row.clash_type for row in text_rows})
        type_mismatch = bool(planner_rows and text_rows and planner_types != text_types)
        notes: list[str] = []
        if type_mismatch:
            notes.append(
                "Sources disagree on clash type: planner="
                f"{','.join(planner_types)}; text={','.join(text_types)}."
            )
        if agreement != "both":
            notes.append("Pair appears in only one source; retained under the conservative union rule.")
        source_ids = sorted({row.source_id for row in rows})
        locations = ";".join(
            f"{source}:{sheet}"
            for source, sheet in sorted({(row.source_id, row.sheet) for row in rows})
        )
        rule = {
            "rule_id": f"CLASH_{repeating}_{target}",
            "scope": scope,
            "rule_type": "timetable_clash",
            "repeating_course": repeating,
            "target_course_code": target,
            "clash_type": _selected_type(rows),
            "source_ids": ";".join(source_ids),
            "sheet": locations,
            "agreement": agreement,
            "notes": " ".join(notes),
        }
        rules.append(rule)
        if agreement != "both" or type_mismatch:
            disagreements.append(
                {
                    **rule,
                    "planner_clash_type": ";".join(planner_types),
                    "text_clash_type": ";".join(text_types),
                    "disagreement": "type_mismatch" if type_mismatch else agreement,
                }
            )
    return rules, disagreements


def _write_csv(path: Path, rows: list[dict[str, str]], fields: tuple[str, ...]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(output, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def build() -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    planner = extract_planner_clashes()
    assert_planner_matches_grid_rules(planner)
    text = extract_text_clashes()
    rules, disagreements = combine_clashes(planner, text)
    fields = (
        "rule_id",
        "scope",
        "rule_type",
        "repeating_course",
        "target_course_code",
        "clash_type",
        "source_ids",
        "sheet",
        "agreement",
        "notes",
    )
    _write_csv(OUTPUT_PATH, rules, fields)
    _write_csv(
        DISAGREEMENTS_PATH,
        disagreements,
        (*fields, "planner_clash_type", "text_clash_type", "disagreement"),
    )
    representatives = {
        clash_type: next(row for row in rules if row["clash_type"] == clash_type)
        for clash_type in CLASH_TYPES
    }
    template_rows = [
        {
            "case_id": f"CLASH_REVIEW_{index:03d}",
            "scope": row["scope"],
            "repeating_modules_json": json.dumps([row["repeating_course"]]),
            "target_course_code": row["target_course_code"],
            "scenario_type": f"{row['clash_type'].lower()}_clash",
            "source_rule_id": row["rule_id"],
            "reference_label": "",
            "expected_material_reasons_json": "",
            "reviewer_ids": "",
            "approval_state": "",
            "reference_authority": "",
            "notes": "Not part of the frozen or scored eligibility set.",
        }
        for index, row in enumerate(representatives.values(), start=1)
    ]
    template_rows.append(
        {
            "case_id": "CLASH_REVIEW_005",
            "scope": "EFA03",
            "repeating_modules_json": json.dumps(["ELEN2003A"]),
            "target_course_code": "ELEN3002A",
            "scenario_type": "no_clash_control",
            "source_rule_id": "",
            "reference_label": "",
            "expected_material_reasons_json": "",
            "reviewer_ids": "",
            "approval_state": "",
            "reference_authority": "",
            "notes": "Clash-only control; other eligibility gates must be assessed separately.",
        }
    )
    _write_csv(
        TEMPLATE_PATH,
        template_rows,
        (
            "case_id",
            "scope",
            "repeating_modules_json",
            "target_course_code",
            "scenario_type",
            "source_rule_id",
            "reference_label",
            "expected_material_reasons_json",
            "reviewer_ids",
            "approval_state",
            "reference_authority",
            "notes",
        ),
    )
    print(f"Planner clash pairs: {len(planner)}")
    print(f"Text clash pairs: {len(text)}")
    print(f"Combined distinct rules: {len(rules)}")
    print(f"Agreement: {dict(sorted(Counter(row['agreement'] for row in rules).items()))}")
    print(f"Disagreements: {len(disagreements)}")
    return rules, disagreements


if __name__ == "__main__":
    build()
