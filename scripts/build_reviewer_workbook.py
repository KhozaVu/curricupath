"""Build the blinded Word workbook used for academic eligibility review."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PACK_DIRECTORY = PROJECT_ROOT / "experiments" / "evaluation" / "v2_5" / "reviewer_pack"
CASE_FILE = PACK_DIRECTORY / "eligibility_cases_blinded.csv"
OUTPUT_FILE = PACK_DIRECTORY / "CurricuPath_Academic_Eligibility_Review_Workbook.docx"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GREY = "F2F4F6"
MID_GREY = "D9D9D9"


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_border(cell, color: str = MID_GREY) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def _set_cell_margins(cell, top: int = 100, start: int = 110, bottom: int = 100, end: int = 110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _format_cell(cell, fill: str | None = None, bold: bool = False) -> None:
    if fill:
        _set_cell_shading(cell, fill)
    _set_cell_border(cell)
    _set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(9)
            run.bold = bold


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    run.font.name = "Aptos Display" if level == 1 else "Aptos"
    run.font.color.rgb = RGBColor(0, 0, 0)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)


def _add_body(document: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        paragraph.add_run(text[len(bold_lead) :])
    else:
        paragraph.add_run(text)
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)


def _add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(10)


def _format_profile(value: str) -> str:
    try:
        profile = json.loads(value)
    except json.JSONDecodeError:
        return value

    def listed(name: str) -> str:
        items = profile.get(name, [])
        return ", ".join(items) if items else "None"

    return "\n".join(
        (
            f"Programme: {profile.get('programme_code', 'Not declared')}",
            f"Completed modules: {listed('completed_modules')}",
            f"Current modules: {listed('current_modules')}",
            f"Planned modules: {listed('planned_modules')}",
        )
    )


def _add_label_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.7)
    table.columns[1].width = Inches(5.55)
    entries = (
        ("Eligible", "The module may be taken under the declared facts and verified evidence."),
        ("Ineligible", "A verified programme, prerequisite, corequisite, or exclusion rule prevents it."),
        ("Manual review", "Evidence is unsuitable or the case requires academic discretion."),
    )
    for index, (label, description) in enumerate(entries):
        row = table.rows[0] if index == 0 else table.add_row()
        row.cells[0].text = label
        row.cells[1].text = description
        _format_cell(row.cells[0], PALE_BLUE, bold=True)
        _format_cell(row.cells[1])


def _add_case(document: Document, case: dict[str, str], index: int) -> None:
    if index > 1:
        document.add_page_break()

    _add_heading(document, f"Case {case['case_id']}")
    _add_body(document, "Review this case independently from the handbook and supplied curriculum evidence. Do not use the production app or ask for a system answer before making your decision.")

    details = document.add_table(rows=0, cols=2)
    details.autofit = False
    details.columns[0].width = Inches(1.55)
    details.columns[1].width = Inches(5.7)
    for label, value in (
        ("Question", case["question"]),
        ("Target module", case["target_course_code"]),
        ("Declared profile", _format_profile(case["student_profile_json"])),
    ):
        cells = details.add_row().cells
        cells[0].text = label
        cells[1].text = value
        _format_cell(cells[0], PALE_BLUE, bold=True)
        _format_cell(cells[1])

    if case.get("clarification_permitted", "false").lower() == "true":
        cells = details.add_row().cells
        cells[0].text = "Clarification"
        cells[1].text = (
            "Use this neutral follow-up fact before deciding the final label:\n"
            + (case.get("scripted_reply") or "No reply supplied.")
        )
        _format_cell(cells[0], PALE_BLUE, bold=True)
        _format_cell(cells[1])

    _add_heading(document, "Reviewer decision", level=2)
    decision = document.add_table(rows=2, cols=3)
    decision.autofit = False
    for cell, label in zip(
        decision.rows[0].cells,
        ("[ ] Eligible", "[ ] Ineligible", "[ ] Manual review required"),
    ):
        cell.text = label
        _format_cell(cell, NAVY, bold=True)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)
    for cell, label in zip(
        decision.rows[1].cells,
        ("[ ] Approved", "[ ] Needs adjudication", "[ ] Case unusable"),
    ):
        cell.text = label
        _format_cell(cell, LIGHT_GREY, bold=True)

    _add_heading(document, "Source backed reasons", level=2)
    _add_body(document, "List the material rule, source ID, and reason for your decision.")
    reasons = document.add_table(rows=4, cols=1)
    for row in reasons.rows:
        row.cells[0].text = "\n"
        _format_cell(row.cells[0])

    _add_heading(document, "Reviewer record", level=2)
    record = document.add_table(rows=2, cols=2)
    for row, values in zip(
        record.rows,
        (("Reviewer ID", "Date"), ("Conflict or clarification note", "Adjudication reference")),
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value + ":\n\n"
            _format_cell(cell, LIGHT_GREY if value in {"Reviewer ID", "Date"} else None)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("CurricuPath academic eligibility review workbook")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 90, 90)


def build_workbook() -> None:
    with CASE_FILE.open(newline="", encoding="utf-8-sig") as csv_file:
        cases = list(csv.DictReader(csv_file))

    document = Document()
    _configure_document(document)
    document.core_properties.title = "CurricuPath Academic Eligibility Review Workbook"
    document.core_properties.subject = "Blinded academic review of synthetic curriculum eligibility cases"
    document.core_properties.author = "CurricuPath"

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("CurricuPath Academic Eligibility Review Workbook")
    subtitle = document.add_paragraph()
    subtitle.add_run("Blinded reviewer copy | EFA03 and EFA04 curriculum cases").italic = True
    subtitle.paragraph_format.space_after = Pt(14)

    _add_body(document, "Please use this workbook to independently approve, reject, or refer synthetic academic eligibility cases for manual review. Your decisions establish the source-backed reference labels used to evaluate CurricuPath. The workbook intentionally does not show any system answer or developer expected label.")

    _add_heading(document, "Reviewer details")
    reviewer = document.add_table(rows=3, cols=2)
    for row, labels in zip(
        reviewer.rows,
        (("Reviewer name or ID", "Affiliation"), ("Role", "Agreement date"), ("Confirmed review hours", "Declared conflict of interest")),
    ):
        for cell, label in zip(row.cells, labels):
            cell.text = label + ":\n\n"
            _format_cell(cell, LIGHT_GREY)

    _add_heading(document, "How to review")
    for item in (
        "Check programme membership first, then prerequisites, corequisites, exclusions, source status, and any academic discretion.",
        "For a clarification-permitted case, use the printed neutral reply before making the final decision.",
        "Use manual review only when verified institutional evidence cannot settle the case or a discretionary academic decision is required.",
        "Write a source-backed reason and your reviewer ID for every decision.",
        "If two reviewers disagree, retain both records and refer the case for adjudication rather than overwriting it.",
    ):
        _add_bullet(document, item)

    _add_heading(document, "Decision labels")
    _add_label_table(document)

    document.add_page_break()
    for index, case in enumerate(cases, start=1):
        _add_case(document, case, index)

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_FILE)


if __name__ == "__main__":
    build_workbook()
